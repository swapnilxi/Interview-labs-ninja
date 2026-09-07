"""Resume → standalone HTML/Markdown/DOCX rendering for export (templates + print/PDF).

One HTML generator serves every export path so there's a single source of truth:
  • browser "Save as PDF" (client prints this HTML),
  • standalone .html download,
  • server-side PDF (xhtml2pdf renders this same HTML).

Because xhtml2pdf supports only a subset of CSS (no flexbox/grid), templates use
print-friendly block + table layouts that also render cleanly in a browser.
"""

from __future__ import annotations

from typing import Any

from ..shared.render import _ACCENTS, _e, _links, _md_escape
from .templates import REGISTRY as _TEMPLATE_REGISTRY
from .templates import get_css as _get_template_css

TEMPLATES = list(_TEMPLATE_REGISTRY.keys())


def _custom_field_items(fields: list | None) -> list[str]:
    """Render a list of {label, value} custom fields as HTML fragments (unwrapped)."""
    out = []
    for f in fields or []:
        if not isinstance(f, dict):
            continue
        label = (f.get("label") or "").strip()
        value = (f.get("value") or "").strip()
        if not (label or value):
            continue
        out.append(f"<b>{_e(label)}:</b> {_e(value)}" if label else _e(value))
    return out


def _contact_line(c: dict) -> str:
    parts = [c.get("email"), c.get("phone"), c.get("location")]
    line = " · ".join(_e(p) for p in parts if p)
    links = _links(c.get("links", []))
    customs = " · ".join(_custom_field_items(c.get("custom_fields")))
    return " · ".join([x for x in [line, links, customs] if x])


def _render_section(section: dict) -> str:
    stype = section.get("section_type")
    content = section.get("content") or {}
    title = section.get("title") or stype
    if stype == "personal_info":
        return ""  # rendered in the header

    body_parts = []
    if stype == "summary":
        text = content.get("text", "")
        if text:
            body_parts.append(f'<p class="summary">{_e(text)}</p>')
    elif stype == "skills":
        groups = content.get("groups", []) or []
        rows = []
        for g in groups:
            name = g.get("name", "")
            items = ", ".join(_e(i) for i in (g.get("items", []) or []))
            if not (name or items):
                continue
            rows.append(f'<p class="skill"><b>{_e(name)}:</b> {items}</p>' if name else f'<p class="skill">{items}</p>')
        extra = content.get("text")
        if extra:
            rows.append(f'<p class="skill">{_e(extra)}</p>')
        body_parts.extend(rows)
    else:
        items = content.get("items", []) or []
        blocks = []
        for it in items:
            it_title = it.get("title", "")
            subtitle = it.get("subtitle", "")
            date = it.get("date", "")
            bullets = it.get("bullets", []) or []
            custom_bits = _custom_field_items(it.get("custom_fields"))
            head = f'<div class="item-head"><span class="item-title">{_e(it_title)}</span>'
            head += f'<span class="item-date">{_e(date)}</span></div>' if date else "</div>"
            sub = f'<div class="item-sub">{_e(subtitle)}</div>' if subtitle else ""
            li_html = "".join(f"<li>{_e(b)}</li>" for b in bullets if b) + "".join(f"<li>{c}</li>" for c in custom_bits)
            blist = f"<ul>{li_html}</ul>" if li_html else ""
            if it_title or subtitle or blist:
                blocks.append(f'<div class="item">{head}{sub}{blist}</div>')
        note = content.get("text")
        if note:
            blocks.insert(0, f'<p class="summary">{_e(note)}</p>')
        body_parts.extend(blocks)

    section_customs = _custom_field_items(content.get("custom_fields"))
    if section_customs:
        body_parts.append('<p class="custom-field">' + " · ".join(section_customs) + "</p>")

    if not body_parts:
        return ""
    return f'<section class="sec"><h2>{_e(title)}</h2>{"".join(body_parts)}</section>'


_BASE_CSS = """
@page { size: A4; margin: 1.6cm; }
* { box-sizing: border-box; }
html { color: #1a1a1a; line-height: 1.4; }
a { color: inherit; text-decoration: none; }
.header .name { margin: 0; }
.header .role { margin: 2px 0 4px; }
.header .contact { font-size: 9.5pt; color: #555; }
.sec { margin-top: 14px; }
.sec h2 { text-transform: uppercase; letter-spacing: 0.06em; }
.item { margin-bottom: 8px; }
.item-head { }
.item-title { font-weight: bold; }
.item-date { color: #666; font-size: 9.5pt; float: right; }
.item-sub { color: #444; font-style: italic; font-size: 10pt; margin-bottom: 2px; }
ul { margin: 3px 0 0 16px; padding: 0; }
li { margin-bottom: 2px; font-size: 10pt; }
.summary, .skill, .custom-field { font-size: 10pt; margin: 2px 0; }
"""


_FONT_STACKS = {
    "serif": "Georgia, 'Times New Roman', serif",
    "sans": "Helvetica, Arial, sans-serif",
    "mono": "'Courier New', monospace",
}


def resume_css_from_spec(spec: dict, accent_hex: str) -> str:
    """Turn a Template-Designer resume spec (see templates_designer/presets.py)
    into print-safe CSS. Layered AFTER _BASE_CSS, so it only sets the knobs it owns."""
    body_font = _FONT_STACKS.get(spec.get("font", "sans"), _FONT_STACKS["sans"])
    name_font = _FONT_STACKS.get(spec.get("nameFont") or spec.get("font", "sans"), body_font)
    density = spec.get("density", "normal")
    size = {"compact": "9.5pt", "normal": "10.5pt", "relaxed": "11pt"}.get(density, "10.5pt")
    sec_gap = {"compact": "9px", "normal": "14px", "relaxed": "18px"}.get(density, "14px")
    page_margin = {"compact": "1.1cm", "normal": "1.6cm", "relaxed": "1.9cm"}.get(density, "1.6cm")
    name_size = {"compact": "18pt", "normal": "21pt", "relaxed": "23pt"}.get(density, "21pt")

    align = "center" if spec.get("headerAlign") == "center" else "left"
    hstyle = spec.get("headerStyle", "plain")
    ink = "#1a1a1a"
    name_color = accent_hex if spec.get("nameColor") == "accent" else ink
    role_color, contact_color = "#444", "#555"

    if hstyle == "rule":
        header_css = f".header {{ text-align:{align}; border-bottom:2px solid #222; padding-bottom:8px; }}"
    elif hstyle == "sidebar":
        header_css = f".header {{ text-align:left; border-left:5px solid {accent_hex}; padding:4px 0 6px 12px; }}"
    elif hstyle == "band":
        header_css = f".header {{ text-align:{align}; background:{accent_hex}; padding:16px 18px; }}"
        name_color, role_color, contact_color = "#fff", "#f4f0ff", "#ece7fb"
    else:  # plain
        header_css = f".header {{ text-align:{align}; padding-bottom:8px; }}"

    h = spec.get("heading", {}) or {}
    hcolor = {"accent": accent_hex, "muted": "#9a9aa5", "ink": "#111"}.get(h.get("color", "ink"), "#111")
    halign = "center" if h.get("align") == "center" else "left"
    hfont = _FONT_STACKS.get(h.get("font"), body_font) if h.get("font") else body_font
    parts = [f"font-size:11pt", f"color:{hcolor}", f"text-align:{halign}", f"font-family:{hfont}"]
    if h.get("uppercase", True):
        spacing = "0.18em" if h.get("spacing") == "wide" else "0.06em"
        parts += ["text-transform:uppercase", f"letter-spacing:{spacing}"]
    else:
        parts += ["text-transform:none", "letter-spacing:0"]
    rule = h.get("rule", "none")
    if rule == "under":
        parts += ["border-bottom:1px solid #bbb", "padding-bottom:2px"]
    elif rule == "leftbar":
        parts += [f"border-left:3px solid {accent_hex}", "padding-left:6px"]

    return (
        f"@page {{ margin:{page_margin}; }}"
        f"body {{ font-family:{body_font}; font-size:{size}; }}"
        f"{header_css}"
        f".header .name {{ font-family:{name_font}; font-size:{name_size}; color:{name_color}; }}"
        f".header .role {{ font-size:11.5pt; color:{role_color}; }}"
        f".header .contact {{ color:{contact_color}; }} .header .contact a {{ color:{contact_color}; }}"
        f".sec {{ margin-top:{sec_gap}; }}"
        ".sec h2 { " + ";".join(parts) + "; }"
    )


def render_resume_html(
    resume: dict, template_key: str | None = None, accent: str = "violet", spec: dict | None = None
) -> str:
    if spec:
        accent_hex = _ACCENTS.get(spec.get("accent") or accent, _ACCENTS["violet"])
        template_css = resume_css_from_spec(spec, accent_hex)
    else:
        template = (template_key or resume.get("template_key") or "classic").lower()
        if template not in _TEMPLATE_REGISTRY:
            template = "classic"
        accent_hex = _ACCENTS.get(accent, _ACCENTS["violet"])
        template_css = _get_template_css(template, accent_hex)

    sections = sorted(resume.get("sections", []), key=lambda s: s.get("sort_order", 0))
    personal = next((s for s in sections if s.get("section_type") == "personal_info"), None)
    pc = (personal or {}).get("content", {}) if personal else {}
    name = pc.get("name") or resume.get("title") or "Your Name"
    role = pc.get("title", "")
    contact = _contact_line(pc)

    header = f'<div class="header"><h1 class="name">{_e(name)}</h1>'
    if role:
        header += f'<div class="role">{_e(role)}</div>'
    if contact:
        header += f'<div class="contact">{contact}</div>'
    header += "</div>"

    body = "".join(_render_section(s) for s in sections if not s.get("is_hidden"))
    css = _BASE_CSS + template_css
    title = _e(name)
    return (
        f"<!DOCTYPE html><html><head><meta charset='utf-8'>"
        f"<title>{title}</title><style>{css}</style></head>"
        f"<body>{header}{body}</body></html>"
    )


def _md_custom_fields(fields: list | None) -> list[str]:
    out = []
    for f in fields or []:
        if not isinstance(f, dict):
            continue
        label = (f.get("label") or "").strip()
        value = (f.get("value") or "").strip()
        if not (label or value):
            continue
        out.append(f"**{_md_escape(label)}:** {_md_escape(value)}" if label else _md_escape(value))
    return out


def render_resume_markdown(resume: dict) -> str:
    """Plain-text/Markdown export — no template styling (that's a print/HTML
    concept), just the content in a clean, portable, diffable shape."""
    sections = sorted(resume.get("sections", []), key=lambda s: s.get("sort_order", 0))
    personal = next((s for s in sections if s.get("section_type") == "personal_info"), None)
    pc = (personal or {}).get("content", {}) if personal else {}
    name = pc.get("name") or resume.get("title") or "Your Name"
    role = pc.get("title", "")

    lines = [f"# {_md_escape(name)}"]
    if role:
        lines.append(f"*{_md_escape(role)}*")
    contact_bits = [pc.get("email"), pc.get("phone"), pc.get("location")]
    contact_bits += [l.get("url") or l.get("label") if isinstance(l, dict) else l for l in (pc.get("links") or [])]
    contact = " · ".join([_md_escape(b) for b in contact_bits if b] + _md_custom_fields(pc.get("custom_fields")))
    if contact:
        lines.append(contact)

    for s in sections:
        if s.get("is_hidden") or s.get("section_type") == "personal_info":
            continue
        content = s.get("content") or {}
        title = s.get("title") or s.get("section_type", "Section")
        block = [f"\n## {_md_escape(title)}"]
        stype = s.get("section_type")
        if stype == "summary":
            if content.get("text"):
                block.append(_md_escape(content["text"]))
        elif stype == "skills":
            for g in content.get("groups", []) or []:
                items = ", ".join(_md_escape(i) for i in (g.get("items", []) or []))
                if g.get("name") and items:
                    block.append(f"- **{_md_escape(g['name'])}:** {items}")
                elif items:
                    block.append(f"- {items}")
            if content.get("text"):
                block.append(_md_escape(content["text"]))
        else:
            for it in content.get("items", []) or []:
                head = f"**{_md_escape(it.get('title', ''))}**"
                if it.get("subtitle"):
                    head += f" — {_md_escape(it['subtitle'])}"
                if it.get("date"):
                    head += f"  ({_md_escape(it['date'])})"
                block.append(head)
                for b in it.get("bullets", []) or []:
                    if b:
                        block.append(f"- {_md_escape(b)}")
                for cf in _md_custom_fields(it.get("custom_fields")):
                    block.append(f"- {cf}")
            if content.get("text"):
                block.append(_md_escape(content["text"]))
        for cf in _md_custom_fields(content.get("custom_fields")):
            block.append(cf)
        if len(block) > 1:
            lines.append("\n".join(block))
    return "\n".join(lines).strip() + "\n"


def _custom_field_pairs(fields: list | None) -> list[tuple[str, str]]:
    out = []
    for f in fields or []:
        if not isinstance(f, dict):
            continue
        label = (f.get("label") or "").strip()
        value = (f.get("value") or "").strip()
        if not (label or value):
            continue
        out.append((label, value))
    return out


def render_resume_docx(resume: dict) -> bytes:
    """Word export via python-docx — same content, a plain business-letter
    layout (DOCX has its own styling conventions; this doesn't try to mirror
    the HTML templates' CSS)."""
    import io

    from docx import Document
    from docx.shared import Pt

    sections = sorted(resume.get("sections", []), key=lambda s: s.get("sort_order", 0))
    personal = next((s for s in sections if s.get("section_type") == "personal_info"), None)
    pc = (personal or {}).get("content", {}) if personal else {}
    name = pc.get("name") or resume.get("title") or "Your Name"
    role = pc.get("title", "")

    doc = Document()
    doc.add_heading(name, level=0)
    if role:
        p = doc.add_paragraph(role)
        p.runs[0].italic = True
    contact_bits = [pc.get("email"), pc.get("phone"), pc.get("location")]
    contact_bits += [l.get("url") or l.get("label") if isinstance(l, dict) else l for l in (pc.get("links") or [])]
    contact = " · ".join(str(b) for b in contact_bits if b)
    custom_contact = _custom_field_pairs(pc.get("custom_fields"))
    if contact or custom_contact:
        cp = doc.add_paragraph()
        wrote = False
        if contact:
            cp.add_run(contact)
            wrote = True
        for label, value in custom_contact:
            if wrote:
                cp.add_run(" · ")
            if label:
                cp.add_run(f"{label}: ").bold = True
            cp.add_run(value)
            wrote = True
        for run in cp.runs:
            run.font.size = Pt(9)

    for s in sections:
        if s.get("is_hidden") or s.get("section_type") == "personal_info":
            continue
        content = s.get("content") or {}
        title = s.get("title") or s.get("section_type", "Section")
        stype = s.get("section_type")
        wrote_any = False

        def ensure_heading():
            nonlocal wrote_any
            if not wrote_any:
                doc.add_heading(title, level=1)
                wrote_any = True

        if stype == "summary":
            if content.get("text"):
                ensure_heading()
                doc.add_paragraph(content["text"])
        elif stype == "skills":
            for g in content.get("groups", []) or []:
                items = ", ".join(g.get("items", []) or [])
                if not items:
                    continue
                ensure_heading()
                p = doc.add_paragraph()
                if g.get("name"):
                    p.add_run(f"{g['name']}: ").bold = True
                p.add_run(items)
            if content.get("text"):
                ensure_heading()
                doc.add_paragraph(content["text"])
        else:
            for it in content.get("items", []) or []:
                custom_pairs = _custom_field_pairs(it.get("custom_fields"))
                if not (it.get("title") or it.get("subtitle") or it.get("bullets") or custom_pairs):
                    continue
                ensure_heading()
                p = doc.add_paragraph()
                if it.get("title"):
                    p.add_run(it["title"]).bold = True
                if it.get("subtitle"):
                    p.add_run(f" — {it['subtitle']}")
                if it.get("date"):
                    p.add_run(f"  ({it['date']})").italic = True
                for b in it.get("bullets", []) or []:
                    if b:
                        doc.add_paragraph(b, style="List Bullet")
                for label, value in custom_pairs:
                    bp = doc.add_paragraph(style="List Bullet")
                    if label:
                        bp.add_run(f"{label}: ").bold = True
                    bp.add_run(value)
            if content.get("text"):
                ensure_heading()
                doc.add_paragraph(content["text"])

        section_customs = _custom_field_pairs(content.get("custom_fields"))
        if section_customs:
            ensure_heading()
            for label, value in section_customs:
                sp = doc.add_paragraph()
                if label:
                    sp.add_run(f"{label}: ").bold = True
                sp.add_run(value)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# Re-exported so callers that do `from . import render` (this module) and then
# `render.html_to_pdf(...)` keep working unchanged after the render.py split —
# html_to_pdf itself is generic and lives in shared/render.py.
from ..shared.render import html_to_pdf  # noqa: E402,F401
