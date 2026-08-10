"""Cover letter → HTML/Markdown/DOCX rendering for export.

A cover letter is prose, not a laid-out document — plain business-letter HTML
(title + paragraphs), print-friendly like the resume templates.
"""

from __future__ import annotations

import html

from ..shared.render import _md_escape


def render_cover_letter_html(letter: dict) -> str:
    title = html.escape(letter.get("title") or "Cover Letter")
    paragraphs = "".join(f"<p>{html.escape(p)}</p>" for p in (letter.get("content_text") or "").split("\n\n") if p.strip())
    return f"""<!doctype html><html><head><meta charset="utf-8"><title>{title}</title>
<style>
body {{ font-family: Georgia, 'Times New Roman', serif; color: #1f2937; max-width: 680px; margin: 40px auto; line-height: 1.6; font-size: 14px; }}
h1 {{ font-size: 18px; font-weight: 600; margin-bottom: 24px; }}
p {{ margin: 0 0 14px 0; white-space: pre-wrap; }}
</style></head><body><h1>{title}</h1>{paragraphs}</body></html>"""


def render_cover_letter_markdown(letter: dict) -> str:
    title = _md_escape(letter.get("title") or "Cover Letter")
    body = letter.get("content_text") or ""
    return f"# {title}\n\n{body}\n"


def render_cover_letter_docx(letter: dict) -> bytes:
    import io

    from docx import Document

    doc = Document()
    doc.add_heading(letter.get("title") or "Cover Letter", level=1)
    for para in (letter.get("content_text") or "").split("\n\n"):
        if para.strip():
            doc.add_paragraph(para.strip())
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# Re-exported so callers that do `from . import render` (this module) and then
# `render.html_to_pdf(...)` keep working unchanged after the render.py split —
# html_to_pdf itself is generic and lives in shared/render.py.
from ..shared.render import html_to_pdf  # noqa: E402,F401
