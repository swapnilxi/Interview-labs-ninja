"""Resume → standalone HTML rendering for export (templates + print/PDF).

One HTML generator serves every export path so there's a single source of truth:
  • browser "Save as PDF" (client prints this HTML),
  • standalone .html download,
  • server-side PDF (xhtml2pdf renders this same HTML).

Because xhtml2pdf supports only a subset of CSS (no flexbox/grid), templates use
print-friendly block + table layouts that also render cleanly in a browser.
"""

from __future__ import annotations

import html
from typing import Any

# Accent palette shared with the portfolio theme vocabulary.
_ACCENTS = {
    "violet": "#7c3aed",
    "emerald": "#059669",
    "blue": "#2563eb",
    "rose": "#e11d48",
    "amber": "#d97706",
    "slate": "#334155",
}

TEMPLATES = ["classic", "modern", "compact", "elegant", "executive", "minimalist", "technical"]


def _e(v: Any) -> str:
    return html.escape("" if v is None else str(v))


def _links(items: list) -> str:
    out = []
    for l in items or []:
        if isinstance(l, dict):
            label = l.get("label") or l.get("url") or ""
            url = l.get("url") or ""
        else:
            label = url = str(l)
        if label:
            out.append(_e(label) if not url else f'<a href="{_e(url)}">{_e(label)}</a>')
    return " · ".join(out)


def _contact_line(c: dict) -> str:
    parts = [c.get("email"), c.get("phone"), c.get("location")]
    line = " · ".join(_e(p) for p in parts if p)
    links = _links(c.get("links", []))
    return " · ".join([x for x in [line, links] if x])


def _render_section(section: dict) -> str:
    stype = section.get("section_type")
    content = section.get("content") or {}
    title = section.get("title") or stype
    if stype == "personal_info":
        return ""  # rendered in the header

    body = ""
    if stype == "summary":
        text = content.get("text", "")
        if not text:
            return ""
        body = f'<p class="summary">{_e(text)}</p>'
    elif stype == "skills":
        groups = content.get("groups", []) or []
        rows = []
        for g in groups:
            name = g.get("name", "")
            items = ", ".join(_e(i) for i in (g.get("items", []) or []))
            if not (name or items):
                continue
            rows.append(f'<p class="skill"><b>{_e(name)}:</b> {items}</p>' if name else f'<p class="skill">{items}</p>')
        extra = content.get("text")
        if extra:
            rows.append(f'<p class="skill">{_e(extra)}</p>')
        if not rows:
            return ""
        body = "".join(rows)
    else:
        items = content.get("items", []) or []
        blocks = []
        for it in items:
            it_title = it.get("title", "")
            subtitle = it.get("subtitle", "")
            date = it.get("date", "")
            bullets = it.get("bullets", []) or []
            head = f'<div class="item-head"><span class="item-title">{_e(it_title)}</span>'
            head += f'<span class="item-date">{_e(date)}</span></div>' if date else "</div>"
            sub = f'<div class="item-sub">{_e(subtitle)}</div>' if subtitle else ""
            blist = "".join(f"<li>{_e(b)}</li>" for b in bullets if b)
            blist = f"<ul>{blist}</ul>" if blist else ""
            if it_title or subtitle or blist:
                blocks.append(f'<div class="item">{head}{sub}{blist}</div>')
        note = content.get("text")
        if note:
            blocks.insert(0, f'<p class="summary">{_e(note)}</p>')
        if not blocks:
            return ""
        body = "".join(blocks)

    return f'<section class="sec"><h2>{_e(title)}</h2>{body}</section>'


_BASE_CSS = """
@page { size: A4; margin: 1.6cm; }
* { box-sizing: border-box; }
html { color: #1a1a1a; line-height: 1.4; }
a { color: inherit; text-decoration: none; }
.header .name { margin: 0; }
.header .role { margin: 2px 0 4px; }
.header .contact { font-size: 9.5pt; color: #555; }
.sec { margin-top: 14px; }
.sec h2 { text-transform: uppercase; letter-spacing: 0.06em; }
.item { margin-bottom: 8px; }
.item-head { }
.item-title { font-weight: bold; }
.item-date { color: #666; font-size: 9.5pt; float: right; }
.item-sub { color: #444; font-style: italic; font-size: 10pt; margin-bottom: 2px; }
ul { margin: 3px 0 0 16px; padding: 0; }
li { margin-bottom: 2px; font-size: 10pt; }
.summary, .skill { font-size: 10pt; margin: 2px 0; }
"""

_TEMPLATE_CSS = {
    "classic": """
        body { font-family: 'Times New Roman', Georgia, serif; font-size: 11pt; }
        .header { text-align: center; border-bottom: 2px solid #222; padding-bottom: 8px; }
        .header .name { font-size: 22pt; }
        .header .role { font-size: 12pt; color: #444; }
        .sec h2 { font-size: 11.5pt; color: #222; border-bottom: 1px solid #bbb; padding-bottom: 2px; }
    """,
    "modern": """
        body { font-family: Helvetica, Arial, sans-serif; font-size: 10.5pt; }
        .header { border-left: 5px solid {ACCENT}; padding: 4px 0 6px 12px; }
        .header .name { font-size: 21pt; color: {ACCENT}; }
        .header .role { font-size: 11.5pt; color: #333; }
        .sec h2 { font-size: 11pt; color: {ACCENT}; }
    """,
    "compact": """
        body { font-family: Helvetica, Arial, sans-serif; font-size: 9.5pt; }
        @page { margin: 1.1cm; }
        .header .name { font-size: 18pt; }
        .header .role { font-size: 10.5pt; color: #444; }
        .sec { margin-top: 9px; }
        .sec h2 { font-size: 10pt; color: #111; border-bottom: 1px solid #ddd; padding-bottom: 1px; }
        li, .summary, .skill { font-size: 9.5pt; }
    """,
    "elegant": """
        body { font-family: Georgia, 'Times New Roman', serif; font-size: 10.5pt; color: #2b2b2b; }
        .header { text-align: center; padding-bottom: 6px; }
        .header .name { font-size: 24pt; letter-spacing: 0.04em; font-weight: normal; }
        .header .role { font-size: 11pt; color: {ACCENT}; letter-spacing: 0.08em; text-transform: uppercase; }
        .header .contact { border-top: 1px solid #ccc; border-bottom: 1px solid #ccc; padding: 4px 0; margin-top: 6px; display: inline-block; }
        .sec h2 { font-size: 11pt; color: {ACCENT}; text-align: center; }
    """,
    "executive": """
        body { font-family: Helvetica, Arial, sans-serif; font-size: 10.5pt; }
        .header { background: {ACCENT}; padding: 16px 18px; }
        .header .name { font-size: 22pt; color: #fff; }
        .header .role { font-size: 12pt; color: #f4f0ff; }
        .header .contact { color: #ece7fb; font-size: 9.5pt; }
        .header .contact a { color: #ece7fb; }
        .sec h2 { font-size: 10.5pt; color: {ACCENT}; text-transform: uppercase; letter-spacing: 0.1em; border-bottom: 2px solid {ACCENT}; padding-bottom: 2px; }
    """,
    "minimalist": """
        body { font-family: Helvetica, Arial, sans-serif; font-size: 10.5pt; color: #333; }
        .header { padding-bottom: 10px; }
        .header .name { font-size: 20pt; font-weight: 600; letter-spacing: 0.02em; }
        .header .role { font-size: 11pt; color: #777; }
        .header .contact { color: #888; }
        .sec { margin-top: 16px; }
        .sec h2 { font-size: 9pt; color: #9a9aa5; text-transform: uppercase; letter-spacing: 0.18em; font-weight: 600; }
        .item-sub { color: #666; }
    """,
    "technical": """
        body { font-family: Helvetica, Arial, sans-serif; font-size: 10.5pt; }
        .header { border-bottom: 2px solid {ACCENT}; padding-bottom: 8px; }
        .header .name { font-size: 20pt; font-family: 'Courier New', monospace; color: {ACCENT}; }
        .header .role { font-size: 11pt; font-family: 'Courier New', monospace; color: #444; }
        .sec h2 { font-size: 10.5pt; color: #111; font-family: 'Courier New', monospace; border-left: 3px solid {ACCENT}; padding-left: 6px; text-transform: none; letter-spacing: 0; }
    """,
}


_FONT_STACKS = {
    "serif": "Georgia, 'Times New Roman', serif",
    "sans": "Helvetica, Arial, sans-serif",
    "mono": "'Courier New', monospace",
}


def resume_css_from_spec(spec: dict, accent_hex: str) -> str:
    """Turn a Template-Designer resume spec (see template_presets.py) into
    print-safe CSS. Layered AFTER _BASE_CSS, so it only sets the knobs it owns."""
    body_font = _FONT_STACKS.get(spec.get("font", "sans"), _FONT_STACKS["sans"])
    name_font = _FONT_STACKS.get(spec.get("nameFont") or spec.get("font", "sans"), body_font)
    density = spec.get("density", "normal")
    size = {"compact": "9.5pt", "normal": "10.5pt", "relaxed": "11pt"}.get(density, "10.5pt")
    sec_gap = {"compact": "9px", "normal": "14px", "relaxed": "18px"}.get(density, "14px")
    page_margin = {"compact": "1.1cm", "normal": "1.6cm", "relaxed": "1.9cm"}.get(density, "1.6cm")
    name_size = {"compact": "18pt", "normal": "21pt", "relaxed": "23pt"}.get(density, "21pt")

    align = "center" if spec.get("headerAlign") == "center" else "left"
    hstyle = spec.get("headerStyle", "plain")
    ink = "#1a1a1a"
    name_color = accent_hex if spec.get("nameColor") == "accent" else ink
    role_color, contact_color = "#444", "#555"

    if hstyle == "rule":
        header_css = f".header {{ text-align:{align}; border-bottom:2px solid #222; padding-bottom:8px; }}"
    elif hstyle == "sidebar":
        header_css = f".header {{ text-align:left; border-left:5px solid {accent_hex}; padding:4px 0 6px 12px; }}"
    elif hstyle == "band":
        header_css = f".header {{ text-align:{align}; background:{accent_hex}; padding:16px 18px; }}"
        name_color, role_color, contact_color = "#fff", "#f4f0ff", "#ece7fb"
    else:  # plain
        header_css = f".header {{ text-align:{align}; padding-bottom:8px; }}"

    h = spec.get("heading", {}) or {}
    hcolor = {"accent": accent_hex, "muted": "#9a9aa5", "ink": "#111"}.get(h.get("color", "ink"), "#111")
    halign = "center" if h.get("align") == "center" else "left"
    hfont = _FONT_STACKS.get(h.get("font"), body_font) if h.get("font") else body_font
    parts = [f"font-size:11pt", f"color:{hcolor}", f"text-align:{halign}", f"font-family:{hfont}"]
    if h.get("uppercase", True):
        spacing = "0.18em" if h.get("spacing") == "wide" else "0.06em"
        parts += ["text-transform:uppercase", f"letter-spacing:{spacing}"]
    else:
        parts += ["text-transform:none", "letter-spacing:0"]
    rule = h.get("rule", "none")
    if rule == "under":
        parts += ["border-bottom:1px solid #bbb", "padding-bottom:2px"]
    elif rule == "leftbar":
        parts += [f"border-left:3px solid {accent_hex}", "padding-left:6px"]

    return (
        f"@page {{ margin:{page_margin}; }}"
        f"body {{ font-family:{body_font}; font-size:{size}; }}"
        f"{header_css}"
        f".header .name {{ font-family:{name_font}; font-size:{name_size}; color:{name_color}; }}"
        f".header .role {{ font-size:11.5pt; color:{role_color}; }}"
        f".header .contact {{ color:{contact_color}; }} .header .contact a {{ color:{contact_color}; }}"
        f".sec {{ margin-top:{sec_gap}; }}"
        ".sec h2 { " + ";".join(parts) + "; }"
    )


def render_resume_html(
    resume: dict, template_key: str | None = None, accent: str = "violet", spec: dict | None = None
) -> str:
    if spec:
        accent_hex = _ACCENTS.get(spec.get("accent") or accent, _ACCENTS["violet"])
        template_css = resume_css_from_spec(spec, accent_hex)
    else:
        template = (template_key or resume.get("template_key") or "classic").lower()
        if template not in _TEMPLATE_CSS:
            template = "classic"
        accent_hex = _ACCENTS.get(accent, _ACCENTS["violet"])
        template_css = _TEMPLATE_CSS[template].replace("{ACCENT}", accent_hex)

    sections = sorted(resume.get("sections", []), key=lambda s: s.get("sort_order", 0))
    personal = next((s for s in sections if s.get("section_type") == "personal_info"), None)
    pc = (personal or {}).get("content", {}) if personal else {}
    name = pc.get("name") or resume.get("title") or "Your Name"
    role = pc.get("title", "")
    contact = _contact_line(pc)

    header = f'<div class="header"><h1 class="name">{_e(name)}</h1>'
    if role:
        header += f'<div class="role">{_e(role)}</div>'
    if contact:
        header += f'<div class="contact">{contact}</div>'
    header += "</div>"

    body = "".join(_render_section(s) for s in sections if not s.get("is_hidden"))
    css = _BASE_CSS + template_css
    title = _e(name)
    return (
        f"<!DOCTYPE html><html><head><meta charset='utf-8'>"
        f"<title>{title}</title><style>{css}</style></head>"
        f"<body>{header}{body}</body></html>"
    )


def _md_escape(v: Any) -> str:
    return "" if v is None else str(v).replace("*", "\\*").replace("_", "\\_")


def render_resume_markdown(resume: dict) -> str:
    """Plain-text/Markdown export — no template styling (that's a print/HTML
    concept), just the content in a clean, portable, diffable shape."""
    sections = sorted(resume.get("sections", []), key=lambda s: s.get("sort_order", 0))
    personal = next((s for s in sections if s.get("section_type") == "personal_info"), None)
    pc = (personal or {}).get("content", {}) if personal else {}
    name = pc.get("name") or resume.get("title") or "Your Name"
    role = pc.get("title", "")

    lines = [f"# {_md_escape(name)}"]
    if role:
        lines.append(f"*{_md_escape(role)}*")
    contact_bits = [pc.get("email"), pc.get("phone"), pc.get("location")]
    contact_bits += [l.get("url") or l.get("label") if isinstance(l, dict) else l for l in (pc.get("links") or [])]
    contact = " · ".join(_md_escape(b) for b in contact_bits if b)
    if contact:
        lines.append(contact)

    for s in sections:
        if s.get("is_hidden") or s.get("section_type") == "personal_info":
            continue
        content = s.get("content") or {}
        title = s.get("title") or s.get("section_type", "Section")
        block = [f"\n## {_md_escape(title)}"]
        stype = s.get("section_type")
        if stype == "summary":
            if content.get("text"):
                block.append(_md_escape(content["text"]))
        elif stype == "skills":
            for g in content.get("groups", []) or []:
                items = ", ".join(_md_escape(i) for i in (g.get("items", []) or []))
                if g.get("name") and items:
                    block.append(f"- **{_md_escape(g['name'])}:** {items}")
                elif items:
                    block.append(f"- {items}")
            if content.get("text"):
                block.append(_md_escape(content["text"]))
        else:
            for it in content.get("items", []) or []:
                head = f"**{_md_escape(it.get('title', ''))}**"
                if it.get("subtitle"):
                    head += f" — {_md_escape(it['subtitle'])}"
                if it.get("date"):
                    head += f"  ({_md_escape(it['date'])})"
                block.append(head)
                for b in it.get("bullets", []) or []:
                    if b:
                        block.append(f"- {_md_escape(b)}")
            if content.get("text"):
                block.append(_md_escape(content["text"]))
        if len(block) > 1:
            lines.append("\n".join(block))
    return "\n".join(lines).strip() + "\n"


def render_resume_docx(resume: dict) -> bytes:
    """Word export via python-docx — same content, a plain business-letter
    layout (DOCX has its own styling conventions; this doesn't try to mirror
    the HTML templates' CSS)."""
    import io

    from docx import Document
    from docx.shared import Pt

    sections = sorted(resume.get("sections", []), key=lambda s: s.get("sort_order", 0))
    personal = next((s for s in sections if s.get("section_type") == "personal_info"), None)
    pc = (personal or {}).get("content", {}) if personal else {}
    name = pc.get("name") or resume.get("title") or "Your Name"
    role = pc.get("title", "")

    doc = Document()
    doc.add_heading(name, level=0)
    if role:
        p = doc.add_paragraph(role)
        p.runs[0].italic = True
    contact_bits = [pc.get("email"), pc.get("phone"), pc.get("location")]
    contact_bits += [l.get("url") or l.get("label") if isinstance(l, dict) else l for l in (pc.get("links") or [])]
    contact = " · ".join(str(b) for b in contact_bits if b)
    if contact:
        cp = doc.add_paragraph(contact)
        cp.runs[0].font.size = Pt(9)

    for s in sections:
        if s.get("is_hidden") or s.get("section_type") == "personal_info":
            continue
        content = s.get("content") or {}
        title = s.get("title") or s.get("section_type", "Section")
        stype = s.get("section_type")
        wrote_any = False

        def ensure_heading():
            nonlocal wrote_any
            if not wrote_any:
                doc.add_heading(title, level=1)
                wrote_any = True

        if stype == "summary":
            if content.get("text"):
                ensure_heading()
                doc.add_paragraph(content["text"])
        elif stype == "skills":
            for g in content.get("groups", []) or []:
                items = ", ".join(g.get("items", []) or [])
                if not items:
                    continue
                ensure_heading()
                p = doc.add_paragraph()
                if g.get("name"):
                    p.add_run(f"{g['name']}: ").bold = True
                p.add_run(items)
            if content.get("text"):
                ensure_heading()
                doc.add_paragraph(content["text"])
        else:
            for it in content.get("items", []) or []:
                if not (it.get("title") or it.get("subtitle") or it.get("bullets")):
                    continue
                ensure_heading()
                p = doc.add_paragraph()
                if it.get("title"):
                    p.add_run(it["title"]).bold = True
                if it.get("subtitle"):
                    p.add_run(f" — {it['subtitle']}")
                if it.get("date"):
                    p.add_run(f"  ({it['date']})").italic = True
                for b in it.get("bullets", []) or []:
                    if b:
                        doc.add_paragraph(b, style="List Bullet")
            if content.get("text"):
                ensure_heading()
                doc.add_paragraph(content["text"])

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ── Portfolio ────────────────────────────────────────────────────────────────

def _render_widget(w: dict, accent_hex: str) -> str:
    wtype = w.get("widget_type")
    c = w.get("content") or {}
    heading = w.get("title") or (wtype or "Section").title()

    if wtype == "hero":
        if not (c.get("headline") or c.get("subheadline") or c.get("tagline")):
            return ""
        cta = f'<div class="cta">{_e(c.get("ctaLabel"))}</div>' if c.get("ctaLabel") else ""
        return (
            f'<div class="hero"><h1>{_e(c.get("headline") or "Your Name")}</h1>'
            + (f'<p class="hero-sub">{_e(c.get("subheadline"))}</p>' if c.get("subheadline") else "")
            + (f'<p class="hero-tag">{_e(c.get("tagline"))}</p>' if c.get("tagline") else "")
            + cta + "</div>"
        )

    if wtype in ("about", "custom"):
        if not c.get("text"):
            return ""
        return f'<section class="w"><h2>{_e(heading)}</h2><p class="para">{_e(c.get("text"))}</p></section>'

    if wtype == "skills":
        groups = c.get("groups", []) or []
        chips = [s for g in groups for s in (g.get("items", []) or []) if s]
        if not chips:
            return ""
        inner = "".join(f'<span class="chip">{_e(s)}</span>' for s in chips)
        return f'<section class="w"><h2>{_e(heading)}</h2><div class="chips">{inner}</div></section>'

    if wtype == "stats":
        items = c.get("items", []) or []
        cells = [it for it in items if it.get("value") or it.get("label")]
        if not cells:
            return ""
        tds = "".join(f'<td class="stat"><div class="stat-v">{_e(it.get("value"))}</div><div class="stat-l">{_e(it.get("label"))}</div></td>' for it in cells)
        return f'<section class="w"><table class="stats"><tr>{tds}</tr></table></section>'

    if wtype == "contact":
        bits = [c.get("email"), c.get("phone"), c.get("location")]
        line = " · ".join(_e(b) for b in bits if b)
        links = _links(c.get("links", []))
        joined = " · ".join(x for x in [line, links] if x)
        if not joined:
            return ""
        return f'<section class="w"><h2>{_e(heading)}</h2><p class="para">{joined}</p></section>'

    # items-based (projects, experience, education, gallery, testimonials)
    items = c.get("items", []) or []
    blocks = []
    for it in items:
        t = it.get("title", "")
        sub = it.get("subtitle", "")
        date = it.get("date", "")
        bullets = "".join(f"<li>{_e(b)}</li>" for b in (it.get("bullets", []) or []) if b)
        bullets = f"<ul>{bullets}</ul>" if bullets else ""
        head = f'<div class="p-head"><b>{_e(t)}</b>' + (f' — <span>{_e(sub)}</span>' if sub else "") + (f'<span class="p-date">{_e(date)}</span>' if date else "") + "</div>"
        if t or sub or bullets:
            blocks.append(f'<div class="p-item">{head}{bullets}</div>')
    if not blocks:
        return ""
    return f'<section class="w"><h2>{_e(heading)}</h2>{"".join(blocks)}</section>'


def _portfolio_template_css(template: str, a: str) -> str:
    """Background + card treatment per visual template (mirrors the frontend's
    portfolioTemplates.ts). Pattern backgrounds gracefully drop to a plain fill
    in xhtml2pdf server PDFs, which don't support CSS gradients."""
    shell = "-webkit-print-color-adjust: exact; print-color-adjust: exact;"
    if template == "isometric":
        pattern = (
            f"background-color:#eef0f7;"
            f"background-image:"
            f"linear-gradient(30deg,{a}20 12%,transparent 12.5%,transparent 87%,{a}20 87.5%,{a}20),"
            f"linear-gradient(150deg,{a}20 12%,transparent 12.5%,transparent 87%,{a}20 87.5%,{a}20),"
            f"linear-gradient(30deg,{a}20 12%,transparent 12.5%,transparent 87%,{a}20 87.5%,{a}20),"
            f"linear-gradient(150deg,{a}20 12%,transparent 12.5%,transparent 87%,{a}20 87.5%,{a}20),"
            f"linear-gradient(60deg,{a}12 25%,transparent 25.5%,transparent 75%,{a}12 75%,{a}12),"
            f"linear-gradient(60deg,{a}12 25%,transparent 25.5%,transparent 75%,{a}12 75%,{a}12);"
            f"background-size:42px 73px;background-position:0 0,0 0,21px 36px,21px 36px,0 0,21px 36px;"
        )
        return f".shell{{{shell}{pattern}}} .page{{background:rgba(255,255,255,.9);box-shadow:0 24px 50px -24px rgba(31,36,48,.4);}}"
    if template == "aurora":
        bg = (
            f"background:"
            f"radial-gradient(closest-side at 18% 12%,{a}33,transparent),"
            f"radial-gradient(closest-side at 84% 6%,#f472b636,transparent),"
            f"radial-gradient(closest-side at 60% 100%,{a}26,transparent),#f4f4fb;"
        )
        return f".shell{{{shell}{bg}}} .page{{background:rgba(255,255,255,.85);box-shadow:0 24px 55px -28px {a}66;}}"
    if template == "blueprint":
        bg = (
            f"background-color:#f6f8fc;"
            f"background-image:linear-gradient({a}14 1px,transparent 1px),linear-gradient(90deg,{a}14 1px,transparent 1px);"
            f"background-size:22px 22px;"
        )
        return f".shell{{{shell}{bg}}} .page{{background:#fff;border-top:3px solid {a};box-shadow:0 12px 30px -18px rgba(31,36,48,.35);}}"
    if template == "dots":
        bg = (
            f"background-color:#f4f5f9;"
            f"background-image:radial-gradient({a}26 1.4px,transparent 1.6px);"
            f"background-size:18px 18px;"
        )
        return f".shell{{{shell}{bg}}} .page{{background:#fff;box-shadow:0 14px 34px -20px rgba(31,36,48,.32);}}"
    if template == "mesh":
        bg = (
            f"background:"
            f"radial-gradient(closest-side at 0% 0%,{a}30,transparent),"
            f"radial-gradient(closest-side at 100% 18%,#22d3ee33,transparent),"
            f"radial-gradient(closest-side at 28% 92%,#f59e0b2e,transparent),"
            f"radial-gradient(closest-side at 92% 100%,{a}26,transparent),#eef0f8;"
        )
        return f".shell{{{shell}{bg}}} .page{{background:rgba(255,255,255,.84);box-shadow:0 24px 55px -28px {a}59;}}"
    if template == "carbon":
        bg = (
            f"background-color:#eef1f5;"
            f"background-image:repeating-linear-gradient(45deg,{a}0d 0,{a}0d 1px,transparent 1px,transparent 11px);"
        )
        return f".shell{{{shell}{bg}}} .page{{background:#fff;border:1px solid #e2e6ee;box-shadow:0 10px 28px -18px rgba(31,36,48,.4);}}"
    # minimal (default)
    return f".shell{{{shell}background:#f1f1f4;}} .page{{background:#fff;box-shadow:0 1px 2px rgba(0,0,0,.06);}}"


def render_portfolio_html(portfolio: dict) -> str:
    theme = portfolio.get("theme") or {}
    accent_hex = _ACCENTS.get(theme.get("accent"), _ACCENTS["violet"])
    serif = theme.get("font") == "serif"
    layout = theme.get("layout", "stack")
    template = theme.get("template", "minimal")
    font_stack = "Georgia, 'Times New Roman', serif" if serif else "-apple-system, Helvetica, Arial, sans-serif"
    card = "background:#fff;border:1px solid #ececf1;border-radius:12px;padding:20px 22px;" if layout == "card" else ""
    centered = "text-align:center;" if layout == "centered" else ""

    widgets = sorted([w for w in portfolio.get("widgets", []) if not w.get("is_hidden")], key=lambda w: w.get("sort_order", 0))
    body = "".join(_render_widget(w, accent_hex) for w in widgets) or '<p class="para">This portfolio is empty.</p>'
    title = _e(portfolio.get("title") or "Portfolio")

    css = f"""
    @page {{ size: A4; margin: 1.4cm; }}
    * {{ box-sizing: border-box; }}
    body {{ margin: 0; color: #1f2430; font-family: {font_stack}; line-height: 1.5; }}
    .shell {{ padding: 28px 16px; }}
    {_portfolio_template_css(template, accent_hex)}
    .page {{ max-width: 860px; margin: 0 auto; padding: 28px 24px; border-radius: 16px; }}
    .w {{ {card} margin-bottom: 22px; {centered} }}
    h1 {{ margin: 0; }}
    h2 {{ color: {accent_hex}; font-size: 18px; margin: 0 0 10px; }}
    .hero {{ background: {accent_hex}; color: #fff; border-radius: 14px; padding: 44px 28px; text-align: center; margin-bottom: 24px; }}
    .hero h1 {{ font-size: 30px; }}
    .hero-sub {{ font-size: 17px; opacity: .92; margin: 6px 0 0; }}
    .hero-tag {{ opacity: .85; margin: 10px auto 0; max-width: 520px; }}
    .cta {{ display: inline-block; margin-top: 16px; background: rgba(255,255,255,.2); padding: 8px 16px; border-radius: 8px; font-weight: 600; }}
    .para {{ margin: 0; white-space: pre-wrap; }}
    .chip {{ display: inline-block; background: {accent_hex}1a; color: {accent_hex}; border-radius: 999px; padding: 3px 10px; font-size: 13px; margin: 0 4px 4px 0; }}
    .stats {{ width: 100%; border-collapse: collapse; }}
    .stat {{ text-align: center; }}
    .stat-v {{ font-size: 24px; font-weight: bold; color: {accent_hex}; }}
    .stat-l {{ font-size: 12px; color: #6b7280; }}
    .p-item {{ margin-bottom: 12px; }}
    .p-date {{ color: #6b7280; font-size: 12px; float: right; }}
    ul {{ margin: 6px 0 0 18px; padding: 0; }}
    li {{ margin-bottom: 3px; }}
    a {{ color: {accent_hex}; text-decoration: none; }}
    """
    return (
        f"<!DOCTYPE html><html><head><meta charset='utf-8'>"
        f"<meta name='viewport' content='width=device-width, initial-scale=1'>"
        f"<title>{title}</title><style>{css}</style></head>"
        f"<body class='shell'><div class='page'>{body}</div></body></html>"
    )


def html_to_pdf(html_str: str) -> bytes:
    """Render HTML to PDF bytes with xhtml2pdf (pure-python, no system deps)."""
    import io

    from xhtml2pdf import pisa

    buf = io.BytesIO()
    result = pisa.CreatePDF(html_str, dest=buf, encoding="utf-8")
    if result.err:
        raise RuntimeError("PDF generation failed")
    return buf.getvalue()
